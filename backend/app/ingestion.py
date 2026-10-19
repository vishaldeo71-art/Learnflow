"""Small, metadata-preserving parsers for VTT/SRT, PDF, PPT, Markdown and text."""

import json
import re
import uuid
from pathlib import Path
from typing import Any

from .domain import Chunk
from .embeddings import EmbeddingProvider

TIMESTAMP = re.compile(r"(?P<h>\d{2}):(?P<m>\d{2}):(?P<s>\d{2})[,.](?P<ms>\d{3})")
INLINE_TIMESTAMP = re.compile(r"\[(?P<time>\d{1,2}:\d{2}:\d{2})\]\s*")


def seconds_from_timestamp(value: str) -> int:
    match = TIMESTAMP.search(value)
    if not match:
        raise ValueError(f"Invalid transcript time: {value}")
    return int(match["h"]) * 3600 + int(match["m"]) * 60 + int(match["s"])


def parse_transcript(text: str) -> list[dict[str, Any]]:
    blocks = re.split(r"\n\s*\n", text.strip())
    segments: list[dict[str, Any]] = []
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        time_line = next((line for line in lines if "-->" in line), None)
        if time_line:
            left, right = time_line.split("-->", 1)
            content = " ".join(line for line in lines[lines.index(time_line) + 1:] if not line.startswith("NOTE"))
            if content:
                segments.append({"start": seconds_from_timestamp(left), "end": seconds_from_timestamp(right), "text": content})
        elif lines[0].startswith("{"):
            try:
                parsed = json.loads("\n".join(lines))
                segments.extend(parsed if isinstance(parsed, list) else parsed.get("segments", []))
            except json.JSONDecodeError:
                continue
    if not segments:
        # Some transcript exports flatten every cue onto one line as
        # `[00:00:03] words [00:00:07] more words` instead of valid VTT/SRT.
        # Recover each cue and use the next cue's start as its exclusive end.
        markers = list(INLINE_TIMESTAMP.finditer(text))
        for index, marker in enumerate(markers):
            start = _seconds_from_clock(marker["time"])
            end = _seconds_from_clock(markers[index + 1]["time"]) if index + 1 < len(markers) else start + 5
            content = text[marker.end():markers[index + 1].start() if index + 1 < len(markers) else len(text)]
            content = re.sub(r"\s+", " ", content).strip()
            if content and end > start:
                segments.append({"start": start, "end": end, "text": content})
    return segments


def _seconds_from_clock(value: str) -> int:
    hours, minutes, seconds = (int(part) for part in value.split(":"))
    return hours * 3600 + minutes * 60 + seconds


def _split_text(text: str, max_chars: int = 1100) -> list[str]:
    paragraphs = [item.strip() for item in re.split(r"\n\s*\n", text) if item.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs or [text.strip()]:
        while len(paragraph) > max_chars:
            if current:
                chunks.append(current)
                current = ""
            split = paragraph.rfind(" ", 0, max_chars)
            split = split if split > 0 else max_chars
            chunks.append(paragraph[:split].strip())
            paragraph = paragraph[split:].strip()
        if current and len(current) + len(paragraph) + 1 > max_chars:
            chunks.append(current)
            current = ""
        current = f"{current}\n\n{paragraph}".strip()
    if current:
        chunks.append(current)
    return chunks


def parse_document(filename: str, raw: bytes) -> tuple[str, list[dict[str, Any]]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".vtt", ".srt", ".json"}:
        text = raw.decode("utf-8-sig", errors="replace")
        segments = parse_transcript(text)
        if not segments:
            if suffix == ".json":
                parsed = json.loads(text)
                segments = parsed if isinstance(parsed, list) else parsed.get("segments", [])
            else:
                raise ValueError("No timestamped transcript cues found. Use WebVTT, SRT, or JSON segments.")
        return "video", segments
    if suffix in {".txt", ".md", ".markdown", ".rst"}:
        text = raw.decode("utf-8-sig", errors="replace")
        segments = parse_transcript(text)
        if segments:
            return "video", segments
        return "notes", [{"text": text}]
    if suffix == ".pdf":
        from pypdf import PdfReader
        import io

        reader = PdfReader(io.BytesIO(raw))
        return "pdf", [{"text": page.extract_text() or "", "page_number": index + 1} for index, page in enumerate(reader.pages)]
    if suffix == ".pptx":
        from pptx import Presentation
        import io

        deck = Presentation(io.BytesIO(raw))
        return "ppt", [
            {"text": "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()), "slide_number": index + 1}
            for index, slide in enumerate(deck.slides)
        ]
    raise ValueError("Supported files: VTT, SRT, JSON transcript, PDF, PPTX, TXT, Markdown.")


async def make_chunks(course_id: str, resource_id: str, name: str, source_type: str, rows: list[dict[str, Any]], embedder: EmbeddingProvider) -> list[Chunk]:
    chunks: list[Chunk] = []
    for row in rows:
        text = str(row.get("text", "")).strip()
        if not text:
            continue
        parts = _split_text(text) if source_type != "video" else [text]
        for part in parts:
            chunks.append(Chunk(
                id=str(uuid.uuid4()),
                course_id=course_id,
                resource_id=resource_id,
                resource_name=name,
                source_type=source_type,
                content=part,
                embedding=await embedder.embed(part),
                start_time=int(row["start"]) if row.get("start") is not None else None,
                end_time=int(row["end"]) if row.get("end") is not None else None,
                page_number=int(row["page_number"]) if row.get("page_number") is not None else None,
                slide_number=int(row["slide_number"]) if row.get("slide_number") is not None else None,
            ))
    if not chunks:
        raise ValueError("This file contains no extractable text.")
    return chunks

