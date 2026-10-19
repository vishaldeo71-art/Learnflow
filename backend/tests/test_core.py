import asyncio
import io
import unittest
import uuid

from app.domain import Chunk
from app.embeddings import feature_embedding
from app.ingestion import make_chunks, parse_document, parse_transcript, seconds_from_timestamp
from app.interview import estimate_cost, extract_skills, make_interview, scorecard
from app.retrieval import Retriever, citation_for
from app.storage import MemoryStore


class FakeEmbedder:
    async def embed(self, text):
        return feature_embedding(text)


def run(coro):
    return asyncio.run(coro)


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.store = MemoryStore()
        self.embedder = FakeEmbedder()
        self.course_id = "00000000-0000-4000-8000-000000000001"

    def add_chunk(self, source_type, content, start=None, end=None, page=None, slide=None, name="sample"):
        chunk = Chunk(
            id=str(uuid.uuid4()), course_id=self.course_id, resource_id="resource", resource_name=name,
            source_type=source_type, content=content, embedding=feature_embedding(content),
            start_time=start, end_time=end, page_number=page, slide_number=slide,
        )
        self.store.chunks[self.course_id].append(chunk)
        return chunk

    def test_video_chunks_after_current_time_never_enter_retrieved_context(self):
        self.add_chunk("video", "Embeddings represent text as vectors.", 60, 90)
        self.add_chunk("video", "Reciprocal rank fusion combines ranked lists.", 400, 470)
        retriever = Retriever(self.store, self.embedder)
        matches, _ = run(retriever.retrieve(self.course_id, "What does reciprocal rank fusion do?", 300))
        self.assertFalse(any(chunk.get("end_time", 0) > 300 for chunk in matches))

    def test_future_detector_returns_metadata_but_not_future_content(self):
        self.add_chunk("video", "Reciprocal rank fusion combines ranked lists.", 400, 470)
        retriever = Retriever(self.store, self.embedder)
        matches, future = run(retriever.retrieve(self.course_id, "What does reciprocal rank fusion do?", 300))
        self.assertEqual(matches, [])
        self.assertIsNotNone(future)
        self.assertNotIn("content", future)

    def test_partly_watched_transcript_chunk_stays_hidden_until_its_end(self):
        self.add_chunk("video", "Reranking scores retrieved passages against the question again.", 400, 470)
        retriever = Retriever(self.store, self.embedder)
        partial, future = run(retriever.retrieve(self.course_id, "What is reranking?", 450))
        complete, _ = run(retriever.retrieve(self.course_id, "What is reranking?", 470))
        self.assertEqual(partial, [])
        self.assertIsNotNone(future)
        self.assertTrue(complete)

    def test_supporting_resources_remain_available_and_citations_keep_page(self):
        self.add_chunk("pdf", "Hybrid retrieval combines exact term and meaning search.", page=7, name="slides.pdf")
        retriever = Retriever(self.store, self.embedder)
        matches, _ = run(retriever.retrieve(self.course_id, "How does hybrid retrieval work?", 0))
        self.assertTrue(matches)
        self.assertEqual(citation_for(matches[0])["label"], "slides.pdf · page 7")

    def test_transcript_parser_keeps_seconds_and_text(self):
        cues = parse_transcript("WEBVTT\n\n00:02:03.000 --> 00:02:08.500\nA timestamped cue.")
        self.assertEqual(cues[0]["start"], 123)
        self.assertEqual(cues[0]["end"], 128)
        self.assertEqual(cues[0]["text"], "A timestamped cue.")
        self.assertEqual(seconds_from_timestamp("00:01:02,000"), 62)

    def test_flattened_youtube_caption_export_keeps_clean_cues_and_time_bounds(self):
        kind, cues = parse_document(
            "transcript.txt",
            "[00:00:00] Lions are large cats [00:00:03] They live in groups [00:00:07] Cubs are born blind".encode(),
        )
        self.assertEqual(kind, "video")
        self.assertEqual(cues, [
            {"start": 0, "end": 3, "text": "Lions are large cats"},
            {"start": 3, "end": 7, "text": "They live in groups"},
            {"start": 7, "end": 12, "text": "Cubs are born blind"},
        ])

    def test_pdf_and_pptx_extract_page_and_slide_metadata(self):
        from pypdf import PdfWriter
        from pptx import Presentation
        from pptx.util import Inches

        pdf = PdfWriter()
        pdf.add_blank_page(width=612, height=792)
        pdf_bytes = io.BytesIO()
        pdf.write(pdf_bytes)
        kind, pages = parse_document("handout.pdf", pdf_bytes.getvalue())
        self.assertEqual(kind, "pdf")
        self.assertEqual(pages[0]["page_number"], 1)

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text = "Slide evidence"
        ppt_bytes = io.BytesIO()
        deck.save(ppt_bytes)
        kind, slides = parse_document("lesson.pptx", ppt_bytes.getvalue())
        self.assertEqual(kind, "ppt")
        self.assertEqual(slides[0]["slide_number"], 1)
        self.assertIn("Slide evidence", slides[0]["text"])

    def test_action_context_still_applies_hard_timestamp_filter(self):
        self.add_chunk("video", "The first section explains embeddings.", 20, 80)
        self.add_chunk("video", "The next section explains reranking.", 100, 140)
        covered = run(self.store.covered_chunks(self.course_id, 100))
        self.assertEqual([item["end_time"] for item in covered], [80])

    def test_resume_jd_skill_extraction_and_scorecard_are_explainable(self):
        self.assertEqual(extract_skills("Built a RAG service with Python, FastAPI, and PostgreSQL."), ["fastapi", "postgresql", "python", "rag"])
        interview = make_interview("AI Engineer", "intermediate", "mixed", 10, "Python FastAPI RAG", "Build a vector search API")
        interview.turns = [
            {"role": "candidate", "content": "I built a Python RAG service, tested alternatives, measured impact, and improved latency."},
            {"role": "interviewer", "content": "How did you validate it?"},
            {"role": "candidate", "content": "I compared outcomes and shared the result."},
        ]
        card = scorecard(interview)
        self.assertEqual(len(card["dimensions"]), 4)
        self.assertEqual(card["answers_assessed"], 2)
        self.assertIn("heuristic", card["method"])
        self.assertGreater(card["overall_score"], 0)

    def test_cost_estimates_cover_required_lengths_and_mock_is_zero(self):
        for minutes in (10, 20, 30):
            result = estimate_cost(minutes, "mock")
            self.assertEqual(result["minutes"], minutes)
            self.assertEqual(result["estimated_cost_usd"], 0)

    def test_ingestion_makes_bounded_chunks_with_provenance(self):
        kind, rows = parse_document("notes.md", b"First paragraph.\n\nSecond paragraph.")
        chunks = run(make_chunks(self.course_id, "resource", "notes.md", kind, rows, self.embedder))
        self.assertEqual(kind, "notes")
        self.assertGreaterEqual(len(chunks), 1)
        self.assertEqual(chunks[0].resource_name, "notes.md")


if __name__ == "__main__":
    unittest.main()
